#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 APP故障库.xlsx 各 sheet 提取中文问题，追加进 config/测试题库.xlsx 的「功能测试库」。

目的：扩充批量测试题库覆盖面（测前准备）。

用法:
  python3 scripts/01_爬取FAQ题库.py
  python3 scripts/01_爬取FAQ题库.py --dry-run
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from yard_agent.paths import FAQ_EXCEL, QUESTION_TEMPLATES_XLSX  # noqa: E402

DEFAULT_EXCEL = FAQ_EXCEL
DEFAULT_OUT = QUESTION_TEMPLATES_XLSX
LIBRARY = "功能测试库"

ZH_Q_ALIASES = {
    "zh（q）",
    "zh(q)",
    "zh（q)",
    "zh(q）",
    "zhq",
    "zh",
    "中文问题",
    "中文（q）",
    "问题",
    "问题现象",
}
EN_Q_ALIASES = {
    "en（q）",
    "en(q)",
    "en（q)",
    "en(q）",
    "enq",
    "en",
    "英文问题",
    "english（q）",
    "question",
    "issue",
}


def _norm_header(cell) -> str:
    if cell is None:
        return ""
    s = str(cell).strip().lower().replace(" ", "")
    s = s.replace("(", "（").replace(")", "）")
    return s


def _find_col_indexes(header_row) -> tuple[int | None, int | None]:
    zh_i = en_i = None
    for i, cell in enumerate(header_row or []):
        h = _norm_header(cell)
        if not h:
            continue
        if zh_i is None and (
            h in {"zh（q）", "zhq", "中文问题", "中文（q）", "问题", "问题现象"}
            or h.startswith("zh（q")
        ):
            zh_i = i
        if en_i is None and (
            h in {"en（q）", "enq", "英文问题", "english（q）", "question", "issue"}
            or h.startswith("en（q")
        ):
            en_i = i
    if zh_i is None or en_i is None:
        for i, cell in enumerate(header_row or []):
            h = _norm_header(cell)
            if zh_i is None and h == "zh":
                zh_i = i
            if en_i is None and h == "en":
                en_i = i
    return zh_i, en_i


def _looks_like_question(text: str) -> bool:
    t = text.strip()
    if not t:
        return False
    if t in {
        "问题",
        "答案",
        "Question",
        "Answer",
        "Possible Causes",
        "Solutions",
        "可能原因",
        "解决措施",
    }:
        return False
    if len(t) > 220:
        return False
    if t.startswith(
        ("1、", "1.", "1)", "①", "Recommended", "Tip:", "提示：", "建议", "请", "Place ", "Choose ", "Check ")
    ):
        return False
    if "？" in t or "?" in t:
        return True
    if len(t) <= 100 and re.match(
        r"^(为什么|如何|怎么|什么|能否|可以|是否|为何|哪|"
        r"Why|How|What|Can |Does |Do |Is |Are |Will |Where|When|Which)",
        t,
        flags=re.I,
    ):
        return True
    if len(t) <= 80 and re.match(
        r"^(机器人|机器|充电|清洁|App|APP|Wi-?Fi|网络|"
        r"The robot|Robot |The App|App |Charging|Cleaning)",
        t,
        flags=re.I,
    ):
        return True
    return False


def _clean_question(text) -> str | None:
    if text is None:
        return None
    q = str(text).strip()
    if not q or q.lower() in {"none", "null", "n/a", "-", "—", "／"}:
        return None
    q = re.sub(r"[ \t]+", " ", q)
    q = re.sub(r"\n+", " ", q).strip()
    if not q or not _looks_like_question(q):
        return None
    return q


def extract_questions_from_excel(excel_path: Path, *, skip_sheet_keywords: list[str] | None = None) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("请先安装 openpyxl: pip install openpyxl") from exc

    skip_sheet_keywords = skip_sheet_keywords or ["废弃", "存稿"]
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    per_sheet: dict[str, dict[str, list[str]]] = {}
    all_zh: list[str] = []
    all_en: list[str] = []

    try:
        for sheet_name in wb.sheetnames:
            if any(k in sheet_name for k in skip_sheet_keywords):
                print(f"跳过 sheet: {sheet_name}")
                continue
            ws = wb[sheet_name]
            rows = ws.iter_rows(values_only=True)
            try:
                header = next(rows)
            except StopIteration:
                continue
            zh_i, en_i = _find_col_indexes(header)
            if zh_i is None and en_i is None:
                print(f"未找到 zh（Q）/en（Q）列，跳过: {sheet_name}")
                print(f"  表头: {header[:10] if header else header}")
                continue

            peek_rows = []
            for _ in range(2):
                try:
                    peek_rows.append(next(rows))
                except StopIteration:
                    break
            data_rows = []
            if peek_rows:
                sub = peek_rows[0]
                sub_vals = [str(c).strip() if c is not None else "" for c in (sub or [])]
                if any(v in {"问题", "答案", "Question", "Answer"} for v in sub_vals):
                    data_rows.extend(peek_rows[1:])
                else:
                    data_rows.extend(peek_rows)
            data_rows.extend(list(rows))

            zh_list: list[str] = []
            en_list: list[str] = []
            for row in data_rows:
                if not row:
                    continue
                if zh_i is not None and zh_i < len(row):
                    q = _clean_question(row[zh_i])
                    if q:
                        zh_list.append(q)
                if en_i is not None and en_i < len(row):
                    q = _clean_question(row[en_i])
                    if q:
                        en_list.append(q)

            per_sheet[sheet_name] = {"zh": zh_list, "en": en_list}
            all_zh.extend(zh_list)
            all_en.extend(en_list)
            print(f"✓ {sheet_name}: zh={len(zh_list)} en={len(en_list)} (列 zh={zh_i}, en={en_i})")
    finally:
        wb.close()

    return {"per_sheet": per_sheet, "zh": all_zh, "en": all_en, "all": all_zh + all_en}


def merge_into_templates(existing: list, new_questions: list[str]) -> tuple[list[str], int]:
    seen: set[str] = set()
    out: list[str] = []
    for q in existing:
        s = str(q).strip()
        if not s or s in seen:
            continue
        seen.add(s)
        out.append(s)
    added = 0
    for q in new_questions:
        s = str(q).strip()
        if not s or s in seen:
            continue
        seen.add(s)
        out.append(s)
        added += 1
    return out, added


def load_existing_questions(path: Path, library: str = LIBRARY) -> list[str]:
    from openpyxl import load_workbook

    if not path.is_file():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    if library not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[library]
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h or "").strip() for h in rows[0]]
    q_i = headers.index("问题") if "问题" in headers else 2
    out: list[str] = []
    for row in rows[1:]:
        if not row or q_i >= len(row):
            continue
        q = str(row[q_i] or "").strip()
        if q:
            out.append(q)
    return out


def append_questions_to_xlsx(path: Path, questions: list[str], library: str = LIBRARY) -> int:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    if library not in wb.sheetnames:
        raise ValueError(f"{path.name} 没有工作表「{library}」")
    ws = wb[library]
    headers = [str(c.value or "").strip() for c in ws[1]]
    q_i = headers.index("问题") if "问题" in headers else 2
    num_i = headers.index("编号") if "编号" in headers else 1
    existing: set[str] = set()
    max_id = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        if q_i < len(row) and row[q_i]:
            existing.add(str(row[q_i]).strip())
        if num_i < len(row) and row[num_i]:
            m = re.search(r"(\d+)$", str(row[num_i]))
            if m:
                max_id = max(max_id, int(m.group(1)))
    from yard_agent.questions import classify_question, draft_cs_answer

    added = 0
    has_type = "问题类型" in headers
    has_answer = "答案" in headers
    for q in questions:
        if q in existing:
            continue
        max_id += 1
        kind = classify_question(q)
        if has_type and has_answer:
            ws.append(["是", f"F{max_id:04d}", kind, q, draft_cs_answer(q, kind)])
        elif has_type:
            ws.append(["是", f"F{max_id:04d}", kind, q])
        else:
            ws.append(["是", f"F{max_id:04d}", q, draft_cs_answer(q, kind)])
        existing.add(q)
        added += 1
    wb.save(path)
    return added


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="从 APP故障库 提取 FAQ 问题并追加到测试题库.xlsx")
    parser.add_argument("--excel", default=str(DEFAULT_EXCEL), help="故障库 Excel 路径")
    parser.add_argument("--out", default=str(DEFAULT_OUT), help="题库 Excel（默认 config/测试题库.xlsx）")
    parser.add_argument("--dry-run", action="store_true", help="只统计不写文件")
    parser.add_argument("--include-draft", action="store_true", help="包含废弃/存稿 sheet")
    args = parser.parse_args(argv)

    excel_path = Path(args.excel)
    out_path = Path(args.out)
    if not excel_path.is_file():
        print(f"找不到 Excel: {excel_path}", file=sys.stderr)
        return 1

    skip = [] if args.include_draft else ["废弃", "存稿"]
    result = extract_questions_from_excel(excel_path, skip_sheet_keywords=skip)
    unique_new = list(dict.fromkeys(result["zh"]))

    existing = load_existing_questions(out_path)
    _merged, added = merge_into_templates(existing, unique_new)

    print("\n==== 汇总 ====")
    print(f"Excel: {excel_path}")
    print(f"sheet 数: {len(result['per_sheet'])}")
    print(f"提取 zh: {len(result['zh'])}  en: {len(result['en'])}（英文不写入题库）")
    print(f"中文去重: {len(unique_new)}")
    print(f"原题库: {len(existing)}  将新增: {added}  合并后: {len(existing) + added}")

    if args.dry_run:
        print("\n[dry-run] 未写入。")
        return 0

    if not out_path.is_file():
        print(f"找不到题库: {out_path}", file=sys.stderr)
        return 1
    written = append_questions_to_xlsx(out_path, unique_new)
    print(f"\n已追加 {written} 题到 {out_path} / {LIBRARY}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
