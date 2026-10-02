# -*- coding: utf-8 -*-
"""Excel 导出与文件名工具。"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from .models import AgentReply
from .scoring import is_refusal
from .style import is_transient_failure


def classify(reply: AgentReply) -> tuple[str, str]:
    """把一条回复判成 已完成 / 风险 / 失败，并给出原因。

    「资料不足」这类也是 Agent 给出的回复，属于答不上来的风险项，
    和脚本没抓到回复的「失败」要分开看。
    """
    detail = f"{reply.status} · {reply.elapsed_sec}s"
    if is_transient_failure(reply):
        return "失败", f"未取到有效回复（{detail}）"
    if is_refusal(reply.answer):
        return "风险", f"Agent 未给出答案，回复为资料不足类说明（{detail}）"
    return "已完成", detail


def safe_name(text: str) -> str:
    return re.sub(r"[^\w\u4e00-\u9fff\-]+", "_", text).strip("_")[:60] or "item"


# 各列默认宽度，未列出的用 20
_COLUMN_WIDTHS = {
    "序号": 6,
    "结果": 10,
    "类型": 10,
    "相似度": 10,
    "类别": 14,
    "日期": 12,
    "产品": 14,
    "用例": 20,
    "探针": 20,
    "步骤": 18,
    "检查点": 24,
    "原因": 34,
    "问题": 44,
    "答案": 70,
    "基线答案": 60,
    "本次答案": 60,
    "截图": 28,
}

# 所有检查类脚本共用的表头；脚本只要把明细写成这几个字段即可
STANDARD_COLUMNS = [
    ("用例", "id"),
    ("类别", "kind"),
    ("问题", "question"),
    ("答案", "answer"),
    ("结果", "result"),
    ("原因", "reason"),
    ("检查点", "note"),
]

# 批量问答没有「检查点」的概念，这一列换成日期，其余与标准表头一致
BATCH_COLUMNS = STANDARD_COLUMNS[:-1] + [("日期", "date")]

# 结果列配色：通过绿、有问题红、其余（新增/缺失等中性变化）黄
_RESULT_FILLS = {
    "通过": "C6EFCE",
    "改善": "C6EFCE",
    "已完成": "C6EFCE",
    "未通过": "FFC7CE",
    "风险": "FFC7CE",
    "退化": "FFC7CE",
    "失败": "FFC7CE",
}
_RESULT_FONTS = {"C6EFCE": "006100", "FFC7CE": "9C0006", "FFEB9C": "9C6500"}


def _style_result_cell(cell, value: str) -> None:
    from openpyxl.styles import Font, PatternFill

    text = str(value or "").strip()
    if not text:
        return
    color = _RESULT_FILLS.get(text, "FFEB9C")
    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(bold=True, color=_RESULT_FONTS[color])


def export_detail_excel(
    details: list[dict],
    excel_path: Path,
    columns: list[tuple[str, str]],
    *,
    sheet_title: str = "测试记录",
) -> Path:
    """把检查类脚本的明细写成 Excel。

    columns 为 [(表头, details 里的字段名), ...]；
    若明细里带 screenshot，会自动追加「截图」列并嵌入缩略图。
    """
    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise SystemExit("缺少 openpyxl，请 pip install openpyxl") from exc

    excel_path = Path(excel_path)
    excel_path.parent.mkdir(parents=True, exist_ok=True)

    has_shot = any(d.get("screenshot") for d in details)
    headers = ["序号"] + [h for h, _ in columns] + (["截图"] if has_shot else [])

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _COLUMN_WIDTHS.get(h, 20)
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(details) + 1}"

    shot_col = len(headers) if has_shot else None
    result_col = headers.index("结果") + 1 if "结果" in headers else None
    for idx, d in enumerate(details, start=1):
        ws.append([idx] + [d.get(key, "") for _, key in columns] + ([""] if has_shot else []))
        r = ws.max_row
        ws.row_dimensions[r].height = 90
        for col in range(1, len(headers) + 1):
            ws.cell(r, col).alignment = Alignment(vertical="top", wrap_text=True)
        if result_col:
            cell = ws.cell(r, result_col)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            _style_result_cell(cell, cell.value)

        if not has_shot:
            continue
        shot = Path(d.get("screenshot", "")) if d.get("screenshot") else None
        if shot and shot.is_file():
            try:
                img = XLImage(str(shot))
                img.width = 180
                img.height = 100
                ws.add_image(img, f"{get_column_letter(shot_col)}{r}")
                ws.cell(r, shot_col).value = shot.name
            except Exception as e:
                ws.cell(r, shot_col).value = f"{shot.name} (嵌入失败: {e})"
        elif d.get("screenshot"):
            ws.cell(r, shot_col).value = d["screenshot"]

    wb.save(excel_path)
    return excel_path.resolve()


def export_excel(records: list[AgentReply], excel_path: Path) -> Path:
    """批量问答导出，表头与检查类脚本一致（末列为日期而非检查点）。"""
    details = []
    seq: dict[str, int] = {}
    for rec in records:
        result, reason = classify(rec)
        seq[rec.product] = seq.get(rec.product, 0) + 1
        details.append(
            {
                "id": f"{safe_name(rec.product)}-{seq[rec.product]:04d}",
                "kind": rec.product,
                "question": rec.question,
                "answer": rec.answer,
                "result": result,
                "reason": reason,
                "date": rec.date or datetime.now().strftime("%Y/%m/%d"),
                "screenshot": rec.screenshot,
            }
        )
    return export_detail_excel(
        details, excel_path, BATCH_COLUMNS, sheet_title="Agent测试记录"
    )
